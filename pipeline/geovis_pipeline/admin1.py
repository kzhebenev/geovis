"""Регионы внутри стран (admin-1): границы Natural Earth + официальная статистика (Росстат).

Ключ региона — код ISO 3166-2 (RU-MOW, KZ-ALA…). Данные по регионам — в том же формате,
что наборы по странам (временной ряд: values, years, series, min/max на все годы).
"""
from __future__ import annotations

import json
from pathlib import Path

import geopandas as gpd
import openpyxl

ADMIN1_COUNTRIES = ["RUS", "KAZ", "BLR", "KGZ", "ARM", "UZB", "TJK"]

# Исправления Natural Earth 10m admin-1 (проверено по площадям полигонов, 08.10.2026):
CODE_FIXES = {"RU-MOS": "RU-MOW", "RU-MOW": "RU-MOS"}   # город Москва (2 854 км²) и область перепутаны кодами
NAME_FIXES = {"RU-ALT": ("Алтайский край", "Altai Krai"), "RU-AL": ("Республика Алтай", "Altai Republic"),
              # устаревшие названия: Алматы (с 1993), Астана (снова с 2022)
              "KZ-ALA": ("Алматы", "Almaty"), "KZ-AST": ("Астана", "Astana")}
DROP = {"RU-X01~"}                                       # безымянный фрагмент 38 км², данных по нему нет


def build_admin1(shp: Path, out_geojson: Path, tolerance: float = 0.01) -> list[dict]:
    """Границы регионов выбранных стран → GeoJSON (code, iso3, ru, en) + список с точками подписей."""
    d = gpd.read_file(shp)
    d = d[d["adm0_a3"].isin(ADMIN1_COUNTRIES) & ~d["iso_3166_2"].isin(DROP)].copy()
    d["code"] = d["iso_3166_2"].map(lambda c: CODE_FIXES.get(c, c))
    d["ru"] = [NAME_FIXES.get(c, (r if isinstance(r, str) else n, None))[0]
               for c, r, n in zip(d["code"], d["name_ru"], d["name"])]
    d["en"] = [NAME_FIXES[c][1] if c in NAME_FIXES else (e if isinstance(e, str) else n)
               for c, e, n in zip(d["code"], d["name_en"], d["name"])]
    d["iso3"] = d["adm0_a3"]
    d["geometry"] = d.geometry.simplify(tolerance, preserve_topology=True)
    pts = d.geometry.representative_point()
    out = d[["code", "iso3", "ru", "en", "geometry"]]
    out_geojson.parent.mkdir(parents=True, exist_ok=True)
    out.to_file(out_geojson, driver="GeoJSON", COORDINATE_PRECISION=3)
    return [{"code": c, "iso3": i, "ru": r, "en": e, "lon": round(p.x, 3), "lat": round(p.y, 3)}
            for c, i, r, e, p in zip(d["code"], d["iso3"], d["ru"], d["en"], pts)]


# Росстат → ISO 3166-2. Явная таблица: строки «в т.ч.» и «без АО» выбраны так, чтобы регионы на карте
# не перекрывались (Архангельская и Тюменская области — без автономных округов, округа — отдельно).
ROSSTAT_CODES = {
    "Белгородская область": "RU-BEL", "Брянская область": "RU-BRY", "Владимирская область": "RU-VLA",
    "Воронежская область": "RU-VOR", "Ивановская область": "RU-IVA", "Калужская область": "RU-KLU",
    "Костромская область": "RU-KOS", "Курская область": "RU-KRS", "Липецкая область": "RU-LIP",
    "Московская область": "RU-MOS", "Орловская область": "RU-ORL", "Рязанская область": "RU-RYA",
    "Смоленская область": "RU-SMO", "Тамбовская область": "RU-TAM", "Тверская область": "RU-TVE",
    "Тульская область": "RU-TUL", "Ярославская область": "RU-YAR", "г.Москва": "RU-MOW", "г. Москва": "RU-MOW",
    "Республика Карелия": "RU-KR", "Республика Коми": "RU-KO",
    "в т.ч. Ненецкий авт. округ": "RU-NEN", "Ненецкий автономный округ": "RU-NEN",
    "Архангельская область без Ненецкого авт.округа": "RU-ARK",
    "Архангельская область без автономного округа": "RU-ARK",
    "Вологодская область": "RU-VLG", "Калининградская область": "RU-KGD", "Ленинградская область": "RU-LEN",
    "Мурманская область": "RU-MUR", "Новгородская область": "RU-NGR", "Псковская область": "RU-PSK",
    "г.Санкт-Петербург": "RU-SPE", "г. Санкт-Петербург": "RU-SPE",
    "Республика Адыгея": "RU-AD", "Республика Калмыкия": "RU-KL", "Краснодарский край": "RU-KDA",
    "Астраханская область": "RU-AST", "Волгоградская область": "RU-VGG", "Ростовская область": "RU-ROS",
    "Республика Дагестан": "RU-DA", "Республика Ингушетия": "RU-IN", "Кабардино-Балкарская Республика": "RU-KB",
    "Карачаево-Черкесская Республика": "RU-KC", "Республика Северная Осетия-Алания": "RU-SE",
    "Республика Северная Осетия - Алания": "RU-SE", "Чеченская Республика": "RU-CE", "Ставропольский край": "RU-STA",
    "Республика Башкортостан": "RU-BA", "Республика Марий Эл": "RU-ME", "Республика Мордовия": "RU-MO",
    "Республика Татарстан": "RU-TA", "Удмуртская Республика": "RU-UD", "Чувашская Республика": "RU-CU",
    "Пермский край": "RU-PER", "Кировская область": "RU-KIR", "Нижегородская область": "RU-NIZ",
    "Оренбургская область": "RU-ORE", "Пензенская область": "RU-PNZ", "Самарская область": "RU-SAM",
    "Саратовская область": "RU-SAR", "Ульяновская область": "RU-ULY",
    "Курганская область": "RU-KGN", "Свердловская область": "RU-SVE",
    "в т.ч. Ханты-Мансийский автономный округ-Югра": "RU-KHM",
    "Ханты-Мансийский автономный округ - Югра": "RU-KHM", "Ямало-Ненецкий автономный округ": "RU-YAN",
    "Тюменская область (без Ханты-Мансийского авт.округа-Югра и Ямало-Ненецкого авт.округа)": "RU-TYU",
    "Тюменская область без автономных округов": "RU-TYU", "Челябинская область": "RU-CHE",
    "Республика Алтай": "RU-AL", "Республика Тыва": "RU-TY", "Республика Хакасия": "RU-KK",
    "Алтайский край": "RU-ALT", "Красноярский край": "RU-KYA", "Иркутская область": "RU-IRK",
    "Кемеровская область-Кузбасс": "RU-KEM", "Кемеровская область - Кузбасс": "RU-KEM",
    "Новосибирская область": "RU-NVS", "Омская область": "RU-OMS", "Томская область": "RU-TOM",
    "Республика Бурятия": "RU-BU", "Республика Саха (Якутия)": "RU-SA", "Забайкальский край": "RU-ZAB",
    "Камчатский край": "RU-KAM", "Приморский край": "RU-PRI", "Хабаровский край": "RU-KHA",
    "Амурская область": "RU-AMU", "Магаданская область": "RU-MAG", "Сахалинская область": "RU-SAK",
    "Еврейская автономная область": "RU-YEV", "Чукотский автономный округ": "RU-CHU",
    # варианты названий в таблицах населения
    "Архангельская область без Ненецкого автономного округа": "RU-ARK", "Кемеровская область": "RU-KEM",
    "Ханты-Мансийский автономный округ-Югра": "RU-KHM",
    "Ямало-Ненецкий авт. округ": "RU-YAN", "в т.ч. Ханты-Мансийский авт. округ-Югра": "RU-KHM",
    "Тюменская область без Ханты-Мансийского автономного округа-Югры и Ямало-Ненецкого автономного округа": "RU-TYU",
    # Крым и Севастополь: в Natural Earth (де-факто граница, та же, что в фокусе «Россия») это полигоны
    # adm0_a3=RUS с кодами UA-43 / UA-40 — данные Росстата сажаем на них
    "Республика Крым": "UA-43", "г.Севастополь": "UA-40", "г. Севастополь": "UA-40",
}
# Агрегаты и составные строки, которые на карту не идут (иначе двойной счёт)
ROSSTAT_SKIP_PREFIXES = ("Валовой", "Российская Федерация", "Российская Федеpация")
ROSSTAT_SKIP = {"Архангельская область", "Тюменская область"}   # «вместе с АО» — на карте АО отдельно


def _clean(name) -> str:
    return " ".join(str(name).replace("\n", " ").split()).strip()


def _year(v) -> int | None:
    s = str(v).strip()
    return int(s[:4]) if s[:4].isdigit() and 1990 <= int(s[:4]) <= 2100 else None


def read_rosstat_series(xlsx: Path, sheets: list[str], factor: float = 1.0) -> tuple[list[dict], list[str]]:
    """Листы «регион × год» → [{code, year, value}] и список несопоставленных регионов (для отчёта)."""
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    rows, unmatched = [], set()
    for sh in sheets:
        years: list[int | None] = []
        for row in wb[sh].iter_rows(values_only=True):
            if not years and sum(_year(c) is not None for c in row[1:]) >= 3:
                years = [_year(c) for c in row]
                continue
            if not years or not row[0]:
                continue
            name = _clean(row[0])
            if name.startswith(ROSSTAT_SKIP_PREFIXES) or "федеральный округ" in name or name in ROSSTAT_SKIP \
                    or name[:1].isdigit():   # сноски «1) Данные…»
                continue
            code = ROSSTAT_CODES.get(name)
            if code is None:
                unmatched.add(name)
                continue
            for y, v in zip(years, row):
                if y and isinstance(v, (int, float)):
                    rows.append({"code": code, "year": y, "value": float(v) * factor})
    return rows, sorted(unmatched)


def read_rosstat_population(xlsx: Path) -> tuple[list[dict], list[str]]:
    """Предварительная оценка численности на 1 января (лист «Численность-человек», колонка «Все население»)."""
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    year, rows, unmatched = None, [], set()
    for row in ws.iter_rows(values_only=True):
        cells = [_clean(c) for c in row if c is not None]
        if year is None:
            for c in cells:
                if c.startswith("На 1 января") and _year(c.split()[-2]):
                    year = _year(c.split()[-2])
            continue
        if not row[0] or not isinstance(row[1], (int, float)):
            continue
        name = _clean(row[0])
        if name.startswith(ROSSTAT_SKIP_PREFIXES) or "федеральный округ" in name or name in ROSSTAT_SKIP:
            continue
        code = ROSSTAT_CODES.get(name)
        if code is None:
            unmatched.add(name)
            continue
        rows.append({"code": code, "year": year, "value": float(row[1])})
    return rows, sorted(unmatched)


def build_series(rows: list[dict], valid_codes: set[str]) -> dict:
    """[{code, year, value}] → формат сайта/API (как наборы по странам): values — последний год, series, общая шкала."""
    rows = [r for r in rows if r["code"] in valid_codes]
    years = sorted({r["year"] for r in rows})
    series = {str(y): sorted(({"code": r["code"], "value": r["value"]} for r in rows if r["year"] == y),
                             key=lambda x: -x["value"]) for y in years}
    latest: dict[str, dict] = {}
    for y in years:
        for v in series[str(y)]:
            latest[v["code"]] = {**v, "year": y}
    allv = [v["value"] for vs in series.values() for v in vs]
    return {"values": sorted(latest.values(), key=lambda x: -x["value"]), "years": years, "series": series,
            "max_value": max(allv), "min_value": min(allv), "count": len(latest)}


def write_json(obj, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, separators=(",", ":")))
